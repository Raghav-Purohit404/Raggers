import argparse
import hashlib
import logging
import os
import pickle
import shutil
import time
from pathlib import Path
from typing import List

import requests
from bs4 import BeautifulSoup
from langchain.schema import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyMuPDFLoader
from langchain_community.vectorstores import FAISS
from langchain_huggingface import HuggingFaceEmbeddings
from pptx import Presentation  # type: ignore

from runtime_paths import (
    BACKEND_RAG_DATA_DIR,
    DATA_DIR,
    EMBEDDING_MODEL_DIR,
    FAISS_INDEX_DIR,
    LOG_DIR,
    ensure_runtime_environment,
)

ensure_runtime_environment()

LOG_DIR.mkdir(parents=True, exist_ok=True)
logger = logging.getLogger(__name__)
if not logger.handlers:
    handler = logging.FileHandler(LOG_DIR / "backend_ingestion.log", encoding="utf-8")
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(name)s: %(message)s"))
    logger.addHandler(handler)
logger.setLevel(logging.INFO)

HASH_STORE_PATH = DATA_DIR / "indexed_hashes.pkl"
INDEX_PATH = FAISS_INDEX_DIR
DEFAULT_DOC_FOLDER = BACKEND_RAG_DATA_DIR

SUPPORTED_EXTENSIONS = [".pdf", ".txt", ".md", ".csv", ".docx", ".ppt", ".pptx"]
MIN_TOKENS = 5
CHUNK_SIZE = 800
CHUNK_OVERLAP = 120


def hash_content(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_indexed_hashes() -> dict:
    if HASH_STORE_PATH.exists():
        try:
            with HASH_STORE_PATH.open("rb") as handle:
                data = pickle.load(handle)
            if isinstance(data, dict):
                return data
            if isinstance(data, set):
                return {str(item): "" for item in data}
        except Exception:
            logger.exception("Unable to load indexed hash store")
    return {}


def save_indexed_hashes(indexed_hashes: dict) -> None:
    HASH_STORE_PATH.parent.mkdir(parents=True, exist_ok=True)
    with HASH_STORE_PATH.open("wb") as handle:
        pickle.dump(indexed_hashes, handle)


def index_exists(index_path) -> bool:
    path = Path(index_path)
    return (path / "index.faiss").exists() and (path / "index.pkl").exists()


def faiss_size(index) -> int:
    if index is None:
        return 0
    try:
        return len(index.docstore._dict)
    except Exception:
        return 0


def get_embedder():
    model_path = str(EMBEDDING_MODEL_DIR) if EMBEDDING_MODEL_DIR.exists() else "sentence-transformers/all-MiniLM-L6-v2"
    return HuggingFaceEmbeddings(model_name=model_path)


def load_ppt_file(path: str) -> List[Document]:
    prs = Presentation(path)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return [Document(page_content=text, metadata={"source": str(Path(path).resolve()), "ingested_by": "backend"})]


def load_file(path: Path) -> List[Document]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        pages = PyMuPDFLoader(str(path)).load()
    elif ext in [".ppt", ".pptx"]:
        pages = load_ppt_file(str(path))
    elif ext in [".txt", ".md", ".csv"]:
        pages = [Document(page_content=path.read_text(encoding="utf-8", errors="ignore"), metadata={})]
    elif ext == ".docx":
        import docx

        document = docx.Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
        pages = [Document(page_content=text, metadata={})]
    else:
        pages = []

    content_hash = file_hash(path)
    for index, doc in enumerate(pages):
        doc.metadata = doc.metadata or {}
        doc.metadata["source"] = str(path.resolve())
        doc.metadata["filename"] = path.name
        doc.metadata["page"] = index + 1
        doc.metadata["source_type"] = "backend"
        doc.metadata["ingested_by"] = "backend"
        doc.metadata["file_hash"] = content_hash
    return pages


def iter_supported_files(folder: Path):
    if not folder.exists():
        return
    for path in folder.rglob("*"):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield path


def load_files(folder: Path, indexed_hashes: dict, rebuild: bool) -> List[Document]:
    docs = []
    for file in iter_supported_files(folder) or []:
        try:
            key = str(file.resolve())
            current_hash = file_hash(file)
            if not rebuild and indexed_hashes.get(key) == current_hash:
                continue
            pages = load_file(file)
            docs.extend(pages)
            indexed_hashes[key] = current_hash
            logger.info("Loaded %s pages/records from %s", len(pages), file)
        except Exception as exc:
            logger.exception("Failed to load %s: %s", file, exc)
    return docs


def load_web(urls: List[str], indexed_hashes: dict, rebuild: bool) -> List[Document]:
    docs = []
    for url in urls:
        try:
            response = requests.get(url, timeout=10)
            response.raise_for_status()
            soup = BeautifulSoup(response.text, "html.parser")
            for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
                tag.decompose()
            text = soup.get_text(separator="\n")
            cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
            current_hash = hash_content(cleaned)
            key = f"url:{url}"
            if not rebuild and indexed_hashes.get(key) == current_hash:
                logger.info("No change in %s; skipping", url)
                continue
            indexed_hashes[key] = current_hash
            docs.append(Document(page_content=cleaned, metadata={"source": url, "source_type": "backend", "ingested_by": "backend"}))
        except Exception as exc:
            logger.exception("Failed to scrape %s: %s", url, exc)
    return docs


def chunk_documents(docs: List[Document]) -> List[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
    chunks = splitter.split_documents(docs)
    filtered = []
    for index, chunk in enumerate(chunks):
        if len(chunk.page_content.strip().split()) >= MIN_TOKENS:
            chunk.metadata = chunk.metadata or {}
            chunk.metadata["chunk_index"] = index
            chunk.metadata["chunk_hash"] = hash_content(chunk.page_content)
            chunk.metadata["rag_snippet"] = chunk.page_content[:500]
            filtered.append(chunk)
    logger.info("Chunk count: %s", len(filtered))
    logger.info("Embedding count: %s", len(filtered))
    return filtered


def deduplicate_chunks(chunks: List[Document], index=None) -> List[Document]:
    if index is None:
        return chunks
    existing_hashes = set()
    try:
        for doc in index.docstore._dict.values():
            value = doc.metadata.get("chunk_hash") if doc.metadata else None
            existing_hashes.add(value or hash_content(doc.page_content))
    except Exception:
        logger.exception("Unable to inspect existing index for dedupe")
        return chunks
    unique = [chunk for chunk in chunks if chunk.metadata.get("chunk_hash") not in existing_hashes]
    logger.info("Unique chunks after dedupe: %s", len(unique))
    return unique


def update_index(chunks: List[Document], index_path=INDEX_PATH, rebuild: bool = False):
    index_path = Path(index_path)
    logger.info("Active FAISS path: %s", index_path)
    embedder = get_embedder()

    existing = None
    if index_exists(index_path) and not rebuild:
        existing = FAISS.load_local(str(index_path), embedder, allow_dangerous_deserialization=True)

    logger.info("FAISS size before ingestion: %s", faiss_size(existing))

    if rebuild:
        if index_path.exists():
            shutil.rmtree(index_path)
        index_path.mkdir(parents=True, exist_ok=True)
        if not chunks:
            logger.info("No chunks available; rebuilt index directory left empty at %s", index_path)
            return None
        index = FAISS.from_documents(chunks, embedder)
    elif existing is not None:
        unique = deduplicate_chunks(chunks, existing)
        index = existing
        if unique:
            index.add_documents(unique)
    else:
        index_path.mkdir(parents=True, exist_ok=True)
        if not chunks:
            logger.info("No chunks available for new index at %s", index_path)
            return None
        index = FAISS.from_documents(chunks, embedder)

    index.save_local(str(index_path))
    logger.info("FAISS size after ingestion: %s", faiss_size(index))
    return index


def run_background_ingestion(
    pdf_dir: Path = DEFAULT_DOC_FOLDER,
    urls: List[str] = None,
    index_path=INDEX_PATH,
    benchmark=False,
    rebuild: bool = False,
):
    if urls is None:
        urls = []

    start = time.time()
    pdf_dir = Path(pdf_dir)
    pdf_dir.mkdir(parents=True, exist_ok=True)
    logger.info("Ingesting from folder: %s", pdf_dir)

    indexed_hashes = {} if rebuild else load_indexed_hashes()
    file_docs = load_files(pdf_dir, indexed_hashes, rebuild)
    web_docs = load_web(urls, indexed_hashes, rebuild)
    all_docs = file_docs + web_docs

    if rebuild:
        live_keys = {str(path.resolve()) for path in iter_supported_files(pdf_dir) or []}
        live_keys.update(f"url:{url}" for url in urls)
        indexed_hashes = {key: value for key, value in indexed_hashes.items() if key in live_keys}

    if not all_docs and not rebuild:
        logger.info("No new or changed documents found in %s", pdf_dir)
        return None

    chunks = chunk_documents(all_docs)
    index = update_index(chunks, index_path=index_path, rebuild=rebuild)
    save_indexed_hashes(indexed_hashes)

    if benchmark:
        logger.info("Ingestion completed in %ss", round(time.time() - start, 2))
    return index


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Backend ingestion script for RAG")
    parser.add_argument("--folder", type=str, default=str(DEFAULT_DOC_FOLDER), help="Folder containing docs to ingest")
    parser.add_argument("--update", action="store_true", help="Update the existing FAISS index")
    parser.add_argument("--rebuild", action="store_true", help="Rebuild the FAISS index from current folder contents")
    parser.add_argument("--benchmark", action="store_true", help="Measure ingestion time")
    parser.add_argument("--index", type=str, default=str(INDEX_PATH), help="Path to FAISS index directory")
    args = parser.parse_args()

    run_background_ingestion(
        pdf_dir=Path(args.folder),
        urls=[],
        index_path=args.index,
        benchmark=args.benchmark,
        rebuild=args.rebuild,
    )
